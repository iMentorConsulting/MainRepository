from fastapi import APIRouter, Depends, Query, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from database import get_db
from models import Expense, Unit
from auth_utils import get_tenant
from pydantic import BaseModel
from typing import Optional
from datetime import date
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

router = APIRouter(prefix="/expenses", tags=["expenses"])

DEFAULT_CATEGORIES = [
    "WEBSITE", "ΠΡΟΜΗΘΕΙΕΣ ΚΡΑΤΗΣΕΩΝ", "ΣΥΝΤΗΡΗΣΕΙΣ", "ΣΥΝΤΗΡΗΣΗ ΚΗΠΩΝ",
    "ΚΑΤΑΣΚΕΥΕΣ", "ΝΕΡΟ", "ΛΕΥΚΑ ΕΙΔΗ", "ΛΟΓΙΣΤΙΚΕΣ ΥΠΗΡΕΣΙΕΣ", "INTERNET",
    "ΕΙΔΗ ΚΑΘΑΡΙΣΜΟΥ", "ΕΠΙΠΛΑ", "ΡΕΥΜΑ", "ΜΙΚΡΟΣΥΣΚΕΥΕΣ", "ΕΡΓΑΛΕΙΑ",
    "ΓΡΑΦΕΙΟΚΡΑΤΙΚΑ", "ΑΠΟΛΥΜΑΝΣΕΙΣ", "ΑΣΤΙΚΗ ΕΥΘΥΝΗ", "WELCOME",
    "ΑΝΑΛΩΣΙΜΑ", "ΕΙΔΗ ΜΠΑΝΙΟΥ", "ΚΑΘΑΡΙΣΜΟΙ", "ΠΛΥΣΙΜΟ ΛΙΝΩΝ", "ΦΟΡΟΙ",
]


class ExpenseIn(BaseModel):
    date: date
    category: str
    item: str
    vendor: Optional[str] = None
    amount: float
    unit_id: Optional[int] = None
    unit_type: Optional[str] = None
    notes: Optional[str] = None


def _serialize(e: Expense) -> dict:
    return {
        "id": e.id,
        "date": e.date.isoformat(),
        "month_year": f"{e.date.month:02d} - {e.date.year}",
        "category": e.category,
        "item": e.item,
        "vendor": e.vendor or "",
        "amount": e.amount,
        "unit_id": e.unit_id,
        "unit_type": e.unit_type,
        "notes": e.notes or "",
    }


@router.get("/categories")
def get_categories(db: Session = Depends(get_db), tenant: str = Depends(get_tenant)):
    existing = db.query(Expense.category).filter(Expense.tenant == tenant).distinct().all()
    custom = [r[0] for r in existing if r[0] not in DEFAULT_CATEGORIES]
    return DEFAULT_CATEGORIES + sorted(custom)


@router.get("/unit-types")
def get_unit_types(db: Session = Depends(get_db), tenant: str = Depends(get_tenant)):
    types = db.query(Unit.type).filter(Unit.tenant == tenant).distinct().all()
    return sorted(set(r[0] for r in types if r[0]))


@router.get("/")
def list_expenses(
    from_date: Optional[date] = None,
    to_date: Optional[date] = None,
    category: Optional[str] = None,
    unit_id: Optional[int] = None,
    unit_type: Optional[str] = None,
    db: Session = Depends(get_db),
    tenant: str = Depends(get_tenant),
):
    q = db.query(Expense).filter(Expense.tenant == tenant)
    if from_date:
        q = q.filter(Expense.date >= from_date)
    if to_date:
        q = q.filter(Expense.date <= to_date)
    if category:
        q = q.filter(Expense.category == category)
    if unit_id:
        q = q.filter(Expense.unit_id == unit_id)
    if unit_type:
        q = q.filter(Expense.unit_type == unit_type)

    expenses = q.order_by(Expense.date.desc(), Expense.id.desc()).all()
    total = round(sum(e.amount for e in expenses), 2)
    return {"expenses": [_serialize(e) for e in expenses], "total": total}


@router.post("/")
def create_expense(data: ExpenseIn, db: Session = Depends(get_db), tenant: str = Depends(get_tenant)):
    e = Expense(tenant=tenant, **data.dict())
    db.add(e)
    db.commit()
    db.refresh(e)
    return _serialize(e)


@router.put("/{expense_id}")
def update_expense(expense_id: int, data: ExpenseIn, db: Session = Depends(get_db), tenant: str = Depends(get_tenant)):
    e = db.query(Expense).filter(Expense.id == expense_id, Expense.tenant == tenant).first()
    if not e:
        raise HTTPException(404)
    for k, v in data.dict().items():
        setattr(e, k, v)
    db.commit()
    return _serialize(e)


@router.delete("/{expense_id}")
def delete_expense(expense_id: int, db: Session = Depends(get_db), tenant: str = Depends(get_tenant)):
    e = db.query(Expense).filter(Expense.id == expense_id, Expense.tenant == tenant).first()
    if not e:
        raise HTTPException(404)
    db.delete(e)
    db.commit()
    return {"ok": True}


@router.get("/template/excel")
def download_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Έξοδα"
    headers = ["Ημερομηνία (ΕΕΕΕ-ΜΜ-ΗΗ)", "Κατηγορία", "Περιγραφή", "Προμηθευτής", "Ποσό €", "Μονάδα (προαιρετικό)"]
    ws.append(headers)
    ws.append(["2025-01-15", "ΡΕΥΜΑ", "Λογαριασμός ΔΕΗ Ιανουαρίου", "ΔΕΗ", 150.00, ""])
    ws.append(["2025-01-21", "ΝΕΡΟ", "Λογαριασμός ΔΕΥΑ", "ΔΕΥΑ", 80.00, ""])
    ws.append(["2025-01-28", "ΚΑΘΑΡΙΣΜΟΙ", "Καθαρισμός μετά αναχώρηση", "ΚΑΘΑΡΙΣΤΡΙΑ", 60.00, "VILLA VERDE"])
    hdr_fill = PatternFill("solid", fgColor="1e3a5f")
    hdr_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 22
    ws2 = wb.create_sheet("Κατηγορίες")
    ws2.append(["Διαθέσιμες Κατηγορίες"])
    ws2["A1"].font = Font(bold=True)
    for cat in DEFAULT_CATEGORIES:
        ws2.append([cat])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=expenses_template.xlsx"},
    )


def _parse_date(val):
    from datetime import datetime as _dt, date as _date
    if val is None:
        raise ValueError("Κενή ημερομηνία")
    if isinstance(val, _date) and not isinstance(val, _dt):
        return val
    if isinstance(val, _dt):
        return val.date()
    s = str(val).strip()
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%m/%d/%Y', '%Y/%m/%d', '%d.%m.%Y'):
        try:
            return _dt.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Μη αναγνώσιμη ημερομηνία: {s!r}")


def _parse_amount(val):
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace('€', '').replace(' ', '').replace(',', '.')
    return float(s)


@router.post("/import/excel")
async def import_expenses(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    tenant: str = Depends(get_tenant),
):
    content = await file.read()
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    ws = wb.active
    imported, errors = 0, []
    units = {u.name.upper(): u.id for u in db.query(Unit).filter(Unit.tenant == tenant).all()}

    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        padded = (list(row) + [None] * 6)[:6]
        date_val, category, item, vendor, amount, unit_ref = padded
        if not any(padded):
            continue
        if not category or not item or not amount:
            errors.append(f"Γραμμή {i}: Λείπουν υποχρεωτικά πεδία (Κατηγορία, Περιγραφή, Ποσό)")
            continue
        try:
            parsed_date = _parse_date(date_val) if date_val else date.today()
            parsed_amount = _parse_amount(amount)
            unit_id = None
            if unit_ref:
                unit_id = units.get(str(unit_ref).upper().strip())
            e = Expense(
                tenant=tenant,
                date=parsed_date,
                category=str(category).upper().strip(),
                item=str(item).strip(),
                vendor=str(vendor).strip() if vendor else None,
                amount=parsed_amount,
                unit_id=unit_id,
            )
            db.add(e)
            imported += 1
        except Exception as ex:
            errors.append(f"Γραμμή {i}: {ex}")

    db.commit()
    return {"imported": imported, "errors": errors}


@router.get("/summary")
def expense_summary(
    from_date: date = Query(...),
    to_date: date = Query(...),
    db: Session = Depends(get_db),
    tenant: str = Depends(get_tenant),
):
    expenses = db.query(Expense).filter(
        Expense.tenant == tenant,
        Expense.date >= from_date,
        Expense.date <= to_date,
    ).all()

    by_cat: dict = {}
    for e in expenses:
        by_cat[e.category] = round(by_cat.get(e.category, 0.0) + e.amount, 2)

    return {
        "total": round(sum(e.amount for e in expenses), 2),
        "count": len(expenses),
        "by_category": [
            {"category": k, "amount": v}
            for k, v in sorted(by_cat.items(), key=lambda x: -x[1])
        ],
    }
